import csv
import hashlib
import json
from html.parser import HTMLParser
from urllib.parse import unquote

from openpyxl import load_workbook
import pytest

from predictor.report.run_report import write_discovery_report, write_run_report


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.paths = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.paths.extend(v for k, v in attrs if k == "href")


def hashes(directory):
    return {str(p.relative_to(directory)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in directory.rglob("*") if p.is_file()}


def bundle(run, name="TF", report=True):
    directory = run / "jobs" / name
    directory.mkdir(parents=True)
    dossier = {"tf_id": name, "family": "ArsR/SmtB", "effector": "As(III)",
               "inducers": {"top": "As(III)", "agreement": "high"},
               "operators": {"primary": {"sequence": "ACGT", "score": 12.0,
                              "provenance": {"accession": "NC", "start": 10, "end": 14,
                                             "locality": "proximal"}}}}
    (directory / "dossier.json").write_text(json.dumps(dossier))
    (directory / "bundle_manifest.json").write_text(json.dumps({"status": "complete", "runtime": {"commit": "pinned"}}))
    if report:
        (directory / "REPORT.html").write_text("<html>Existing report</html>")
    return directory


def test_infer_legacy_bundles_links_and_hash_neutrality(tmp_path):
    bundle(tmp_path, "with spaces")
    bundle(tmp_path, "no_report", report=False)
    before = hashes(tmp_path / "jobs")
    discovery = write_discovery_report(tmp_path)
    report = write_run_report(tmp_path)
    assert hashes(tmp_path / "jobs") == before
    with (tmp_path / "Supplementary_Table_S1.tsv").open() as stream:
        rows = list(csv.DictReader(stream, delimiter="\t"))
    assert len(rows) == 2
    assert {r["status"] for r in rows} == {"complete", "bundle present / report missing"}
    for path in (report, discovery):
        parser = Links()
        parser.feed(path.read_text())
        assert all((path.parent / unquote(link)).exists() for link in parser.paths)
    book = load_workbook(tmp_path / "Supplementary_Table_S1.xlsx")
    sheet = book.active
    column = [c.value for c in sheet[1]].index("report_path") + 1
    cell = next(sheet.cell(i, column) for i in range(2, sheet.max_row + 1) if sheet.cell(i, column).hyperlink)
    assert (tmp_path / unquote(cell.hyperlink.target)).is_file()
    assert sheet.freeze_panes == "C2" and sheet.auto_filter.ref
    assert "pinned" in report.read_text()


def test_hidden_publication_drafts_do_not_inflate_stored_run_counts(tmp_path):
    bundle(tmp_path, "legitimate")
    bundle(tmp_path, ".TF.stage-interrupted")
    before = hashes(tmp_path / "jobs")
    discovery = write_discovery_report(tmp_path)
    write_run_report(tmp_path)
    rows = list(csv.DictReader((tmp_path / "Supplementary_Table_S1.tsv").open(), delimiter="\t"))
    assert [row["run_name"] for row in rows] == ["legitimate"]
    assert "<b>1</b> candidate in" in discovery.read_text()
    assert hashes(tmp_path / "jobs") == before


def test_external_bundle_link_not_counted_as_stored_candidate(tmp_path):
    run = tmp_path / "run"
    run.mkdir()
    bundle(run, "legitimate")
    outside = bundle(tmp_path / "outside", "external")
    try:
        (run / "jobs" / "escape").symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip("symlink creation unavailable")
    write_run_report(run)
    rows = list(csv.DictReader((run / "Supplementary_Table_S1.tsv").open(), delimiter="\t"))
    assert [row["run_name"] for row in rows] == ["legitimate"]


def test_manifest_missing_candidate_and_html_escaping(tmp_path):
    with (tmp_path / "manifest.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, ["run_name", "family", "product", "fasta", "discovery_engine"])
        writer.writeheader()
        writer.writerow({"run_name": "missing", "family": "Fur", "product": '<img src=x onerror="attack()">',
                         "fasta": "", "discovery_engine": "provided_protein"})
    discovery = write_discovery_report(tmp_path)
    report = write_run_report(tmp_path)
    assert "<b>1</b> candidate in" in discovery.read_text()
    assert "provided_protein" in discovery.read_text()
    assert "not run / bundle missing" in report.read_text()
    assert "&lt;img" in report.read_text() and '<img src=x' not in report.read_text()
    assert "addEventListener" in report.read_text()


def test_raw_engine_table_preserved(tmp_path):
    discovery = tmp_path / "discovery"
    discovery.mkdir()
    raw = discovery / "candidates.tsv"
    raw.write_text("protein_id\tfamily\tevalue\nP1\tArsR/SmtB\t1e-30\n")
    before = raw.read_bytes()
    target = write_discovery_report(tmp_path)
    assert raw.read_bytes() == before
    assert "1e-30" in target.read_text()
    assert (discovery / "discovery_candidates.tsv").is_file()


def test_formula_labels_stay_literals_and_no_bundle_writes(tmp_path):
    bundle(tmp_path)
    with (tmp_path / "summary.tsv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, ["run_name", "product"], delimiter="\t")
        writer.writeheader()
        writer.writerow({"run_name": "TF", "product": "=HYPERLINK(\"bad\")"})
    write_run_report(tmp_path)
    sheet = load_workbook(tmp_path / "Supplementary_Table_S1.xlsx").active
    col = [c.value for c in sheet[1]].index("product") + 1
    assert sheet.cell(2, col).data_type == "s"
    with pytest.raises(ValueError, match="outside"):
        write_run_report(tmp_path / "jobs" / "TF")
    with pytest.raises(ValueError, match="outside"):
        write_discovery_report(tmp_path / "jobs" / "TF")


def test_candidate_json_evidence_and_summary_extra_columns(tmp_path):
    directory = tmp_path / "discovery"
    directory.mkdir()
    (directory / "candidates.json").write_text(json.dumps([{"protein_id": "P1", "family": "Fur",
                                                            "profile": "ecFur", "evidence": {"evalue": "1e-70"}}]))
    (tmp_path / "protein.faa").write_text(">P1 annotation\nMACK\n")
    with (tmp_path / "manifest.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, ["run_name", "fasta", "family", "discovery_engine", "custom_identifier"])
        writer.writeheader()
        writer.writerow({"run_name": "different", "fasta": "protein.faa", "family": "Fur",
                         "discovery_engine": "bitacora", "custom_identifier": "original-source"})
    discovery = write_discovery_report(tmp_path)
    assert "1e-70" in discovery.read_text() and "ecFur" in discovery.read_text()
    write_run_report(tmp_path)
    with (tmp_path / "Supplementary_Table_S1.tsv").open() as stream:
        rows = list(csv.DictReader(stream, delimiter="\t"))
    assert rows[0]["custom_identifier"] == "original-source"
    with pytest.raises(ValueError, match="does not exist"):
        write_run_report(tmp_path / "missing")


def test_unselected_candidates_are_not_abstentions(tmp_path):
    """A genome run predicting one selected candidate must not report the others as abstentions."""
    bundle(tmp_path, "TF")
    with (tmp_path / "manifest.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, ["run_name", "family", "fasta", "discovery_engine"])
        writer.writeheader()
        writer.writerow({"run_name": "TF", "family": "ArsR", "fasta": "", "discovery_engine": "bitacora"})
        writer.writerow({"run_name": "other", "family": "GntR", "fasta": "", "discovery_engine": "bitacora"})
    with (tmp_path / "summary.tsv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, ["run_name", "family", "status"], delimiter="	")
        writer.writeheader()
        writer.writerow({"run_name": "other", "family": "", "status": "not run"})
    report = write_run_report(tmp_path).read_text()
    with (tmp_path / "Supplementary_Table_S1.tsv").open() as stream:
        rows = {r["run_name"]: r for r in csv.DictReader(stream, delimiter="	")}
    assert rows["other"]["discovery_family"] == "GntR"
    assert rows["other"]["inducer_class"] == "not predicted" and rows["other"]["abstention"] == ""
    assert rows["TF"]["primary_locality"] == "proximal"
    assert "_predicted" not in rows["TF"]
    assert "<strong>1</strong><span>not predicted</span>" in report
    assert "Discovered, not predicted (1)" in report and "GntR" in report
    assert "abstained" not in report


def test_nearest_reference_is_not_presented_as_query_gene(tmp_path):
    directory = bundle(tmp_path, "unannotated_query")
    dossier = json.loads((directory / "dossier.json").read_text())
    dossier["tf_record"] = {"nearest_tf": "REFERENCE_NOT_QUERY"}
    (directory / "dossier.json").write_text(json.dumps(dossier))
    report = write_run_report(tmp_path).read_text()
    assert "unannotated_query" in report
    assert "REFERENCE_NOT_QUERY" not in report


def test_legacy_list_operators_and_invalid_json_shapes_remain_reportable(tmp_path):
    directory = bundle(tmp_path, "legacy")
    dossier = json.loads((directory / "dossier.json").read_text())
    dossier["operators"] = [{"provenance": {"locality": "proximal"}}]
    (directory / "dossier.json").write_text(json.dumps(dossier))
    (directory / "run").mkdir()
    (directory / "run" / "run_record.json").write_text("[]")
    (tmp_path / "run_input.json").write_text("[]")
    invalid = bundle(tmp_path, "invalid")
    (invalid / "dossier.json").write_text("[]")
    report = write_run_report(tmp_path).read_text()
    assert "yes (#1)" in report
    assert "dossier missing or invalid" in report


def test_failed_status_survives_missing_old_report(tmp_path):
    bundle(tmp_path, "failed_resume", report=False)
    (tmp_path / "summary.tsv").write_text("run_name\tstatus\nfailed_resume\tfailed (resume mismatch)\n")
    assert "failed (resume mismatch)" in write_run_report(tmp_path).read_text()


def test_provenance_timestamps_are_converted_to_utc():
    from predictor.report.run_page import _when
    assert _when("2026-10-10T10:30:00-03:00") == "2026-10-10 13:30 UTC"
    assert _when("2026-10-10T10:30:00Z") == "2026-10-10 10:30 UTC"
    assert _when(None) == ""


def test_regulated_gene_display_uses_saved_locus_end_and_strand():
    from predictor.report.run_page import _regulated
    dossier = {"operators": {"primary": {"provenance": {"start": 10, "end": 20, "strand": "+"}}},
               "per_hit_regulation": [
                   {"operator": {"start": 10, "end": 30, "strand": "-"}, "regulated_gene": "wrong_gene"},
                   {"operator": {"start": 10, "end": 20, "strand": "+"}, "regulated_gene": "saved_gene"}]}
    value = _regulated(dossier, 10)
    assert "saved_gene" in value and "wrong_gene" not in value
    assert "saved_gene" not in _regulated(dossier, None)


def test_sortable_family_groups_have_separate_bodies_and_keyboard_headers(tmp_path):
    class GroupedTable(HTMLParser):
        def __init__(self):
            super().__init__()
            self.bodies = []
            self.headers = []
            self.buttons = 0
            self.current = None

        def handle_starttag(self, tag, attrs):
            attrs = dict(attrs)
            if tag == "tbody":
                self.current = []
                self.bodies.append(self.current)
            elif tag == "tr" and self.current is not None:
                self.current.append(attrs.get("class", ""))
            elif tag == "th":
                self.headers.append(attrs)
            elif tag == "button":
                self.buttons += 1

        def handle_endtag(self, tag):
            if tag == "tbody":
                self.current = None

    for name, family in (("second", "Fur"), ("first", "Fur"), ("other", "ArsR/SmtB")):
        directory = bundle(tmp_path, name)
        dossier = json.loads((directory / "dossier.json").read_text())
        dossier["family"] = family
        (directory / "dossier.json").write_text(json.dumps(dossier))
    parser = GroupedTable()
    parser.feed(write_run_report(tmp_path).read_text())
    assert len(parser.bodies) == 2
    assert sorted(len(rows) for rows in parser.bodies) == [2, 3]
    assert all(rows[0] == "group" and rows.count("group") == 1 for rows in parser.bodies)
    assert len(parser.headers) == parser.buttons == 8
    assert all(header["aria-sort"] == "none" and header["scope"] == "col" for header in parser.headers)


def test_confidence_distribution_uses_actual_saved_confidence(tmp_path):
    first = bundle(tmp_path, "recorded")
    bundle(tmp_path, "not_recorded")
    dossier = json.loads((first / "dossier.json").read_text())
    dossier["operators"]["primary"]["confidence"] = "high"
    (first / "dossier.json").write_text(json.dumps(dossier))
    report = write_run_report(tmp_path).read_text()
    assert "Saved family calls" in report and "Saved inducer classes" in report
    assert "<dt>high</dt><dd>1</dd>" in report
    assert "Not recorded for 1 predicted TFs" in report
    with (tmp_path / "Supplementary_Table_S1.tsv").open() as stream:
        rows = {row["run_name"]: row for row in csv.DictReader(stream, delimiter="\t")}
    assert rows["recorded"]["operator_confidence"] == "high"
    assert rows["not_recorded"]["operator_confidence"] == ""
    assert rows["not_recorded"]["primary_locality"] == "proximal"


def test_missing_saved_confidence_is_explicitly_unavailable(tmp_path):
    bundle(tmp_path)
    report = write_run_report(tmp_path).read_text()
    assert "Saved operator confidence" in report
    assert "Unavailable: no saved values in these bundles" in report
    assert "<dt>proximal</dt>" not in report


def test_run_records_are_linked_not_embedded(tmp_path):
    """Each bundle keeps its full run record; copying them all into the page made it ~1 MB."""
    directory = bundle(tmp_path)
    (directory / "run").mkdir()
    (directory / "run" / "run_record.json").write_text(json.dumps({"stages": ["x" * 5000]}))
    report = write_run_report(tmp_path).read_text()
    assert "x" * 5000 not in report
    assert 'href="jobs/TF/run/run_record.json"' in report


def test_discovery_page_links_predictions_and_labels_raw_files(tmp_path):
    bundle(tmp_path, "TF")
    with (tmp_path / "manifest.csv").open("w", newline="") as stream:
        writer = csv.DictWriter(stream, ["run_name", "family", "gene", "fasta", "discovery_engine",
                                         "discovery_profile"])
        writer.writeheader()
        writer.writerow({"run_name": "TF", "family": "ArsR", "gene": "arsR", "fasta": "",
                         "discovery_engine": "bitacora", "discovery_profile": "BsCzrA|mtNmtR"})
        writer.writerow({"run_name": "other", "family": "GntR", "gene": "lldR", "fasta": "",
                         "discovery_engine": "bitacora", "discovery_profile": "GntR_LldR"})
    raw = tmp_path / "discovery" / "BsCzrA"
    raw.mkdir(parents=True)
    (raw / "hits.txt").write_text("")
    page = write_discovery_report(tmp_path).read_text()
    assert 'href="../jobs/TF/REPORT.html"' in page and "not predicted" in page
    assert "BsCzrA, mtNmtR" in page and "(empty: no records)" in page
    assert ">best E-value</button>" not in page


def test_discovery_profile_labels_are_not_filesystem_paths(tmp_path):
    from predictor.report import discovery_page
    run = tmp_path / "run"
    discovery = run / "discovery"
    discovery.mkdir(parents=True)
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "private-file.txt").write_text("private retained content")
    page = discovery_page.render([{"protein_id": '<unsafe"protein>', "family": '<family>',
                                   "profile": "../../outside", "product": '<script>alert(1)</script>'}],
                                 discovery=discovery, reports={})
    assert "private-file.txt" not in page and "private retained content" not in page
    assert "../../outside" in page  # label is retained honestly without treating it as a path
    assert "<script>alert" not in page and "&lt;script&gt;" in page and "&lt;family&gt;" in page
    assert 'href="../RUN_REPORT.html"' not in page


def test_profile_evidence_has_keyboard_sorting_and_grouped_candidate_bodies(tmp_path):
    from predictor.report import discovery_page
    discovery = tmp_path / "discovery"
    (discovery / "profile_one").mkdir(parents=True)
    (discovery / "profile_one" / "raw.tsv").write_text("record")
    page = discovery_page.render([{"protein_id": "one", "family": "ArsR", "profile": "profile_one"},
                                  {"protein_id": "two", "family": "GntR", "profile": "profile_two"}],
                                 discovery=discovery, reports={})
    assert page.count('<tr class="group">') == 2
    assert page.count('<tbody>') == 3  # two family groups plus independent profile table
    assert page.count('<th scope="col" aria-sort="none"><button type="button">') == 9
    assert '<button type="button">profile</button>' in page
    assert '<button type="button">candidates</button>' in page


def test_discovery_and_record_links_cannot_escape_saved_run(tmp_path):
    run = tmp_path / "run"
    run.mkdir()
    directory = bundle(run)
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "REPORT.html").write_text("external report")
    (outside / "run_record.json").write_text(json.dumps({"runtime": {"secret_marker": "external-record"}}))
    (outside / "private-file.txt").write_text("external raw file")
    (directory / "REPORT.html").unlink()
    (directory / "run").mkdir()
    discovery = run / "discovery"
    discovery.mkdir()
    raw = discovery / "inside-profile"
    raw.mkdir()
    try:
        (directory / "REPORT.html").symlink_to(outside / "REPORT.html")
        (directory / "run" / "run_record.json").symlink_to(outside / "run_record.json")
        (discovery / "outside-profile").symlink_to(outside, target_is_directory=True)
        (raw / "linked-private.txt").symlink_to(outside / "private-file.txt")
        (discovery / "loose-private.txt").symlink_to(outside / "private-file.txt")
    except OSError:
        pytest.skip("symlink creation unavailable")
    page = write_discovery_report(run).read_text()
    global_page = write_run_report(run).read_text()
    assert "outside-profile" not in page and "private-file.txt" not in page
    assert "linked-private.txt" not in page and "loose-private.txt" not in page
    assert 'href="../jobs/TF/REPORT.html"' not in page
    assert 'href="jobs/TF/REPORT.html"' not in global_page
    assert 'href="jobs/TF/run/run_record.json"' not in global_page
    assert "external-record" not in global_page


def test_provenance_record_count_matches_available_links(tmp_path):
    first = bundle(tmp_path, "one")
    bundle(tmp_path, "two")
    (first / "run").mkdir()
    (first / "run" / "run_record.json").write_text("{}")
    page = write_run_report(tmp_path).read_text()
    assert "Per-TF run records (1)" in page


def test_scientific_scores_keep_numeric_sort_keys_despite_bits_annotation(tmp_path):
    from predictor.report import discovery_page
    discovery = tmp_path / "discovery"
    discovery.mkdir()
    page = discovery_page.render([{"protein_id": "one", "evalue": "1e-10", "bits": "28"},
                                  {"protein_id": "two", "evalue": "1e-6", "bits": "36"}],
                                 discovery=discovery, reports={})
    assert 'data-sort-value="1e-10"' in page and 'data-sort-value="1e-6"' in page
    assert "28 bits" in page and "36 bits" in page
    from predictor.report.run_report import SORT
    assert "getAttribute('data-sort-value')??cell.textContent.trim()" in SORT
