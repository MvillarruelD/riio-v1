"""build_docx.py -- the run document: one .docx per complete (six-genome) run.

Layout follows the earlier Pipeline_validation_figures.docx: US Letter, 25 mm margins, one section;
title, subtitle and a provenance line; a validation summary; then the main and the supplementary
figures, each as image -> bold "Figure N | title" -> legend. Every figure and legend comes from
build_legends.figure_entries(); every definition from numbers.json. Nothing here holds a result.

    python reports/build_docx.py
"""
from __future__ import annotations

import json
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

from build_legends import figure_entries, provenance_line

HERE = Path(__file__).resolve().parent
FIGDIR = HERE / "figures"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x4A, 0x7A)
GREY = RGBColor(0x5F, 0x6A, 0x66)
PAGE_W, PAGE_H, MARGIN = Mm(215.9), Mm(279.4), Mm(25)
TEXT_W = PAGE_W - 2 * MARGIN
#: the tallest a figure may be drawn, leaving room for its title and a few legend lines on the page
MAX_FIG_H = PAGE_H - 2 * MARGIN - Mm(45)
#: legend type size, and the room the caption line and paragraph spacing take on a page
LEGEND_PT = 8.5
CAPTION_MM = 14.0
#: Arial's average character width in em, and line pitch, for estimating a legend's height
CHAR_EM, LINE_PITCH = 0.5, 1.2


def legend_height(text: str) -> int:
    """Estimated height of a legend paragraph (EMU), so its figure can be sized to share the page."""
    plain = re.sub(r"\*\*|\*", "", text)
    chars_per_line = TEXT_W / Pt(LEGEND_PT * CHAR_EM)
    lines = sum(max(1, -(-len(par) // int(chars_per_line))) for par in plain.split("\n"))
    return int(Pt(LEGEND_PT * LINE_PITCH) * (lines + 0.5))
#: definitions quoted in the summary, in order, with the term shown for each
DEFINITIONS = [("metal_call", "Metal call"), ("ion_resolved", "Ion-resolved"),
               ("operator_supports_call", "Operator supports the call"),
               ("corroborated_metal_call", "Corroborated metal call"),
               ("primary_operator", "Primary operator"), ("regulon_cap", "Regulon cap")]


def run_suffix(N: dict) -> str:
    """'ecoli_rerun20260924', 'salm_rerun20260924', ... -> 'rerun20260924' (the run's shared name)."""
    tails = {t["tag"].split("_", 1)[-1] for t in N["provenance"]["runs"].values()}
    return tails.pop() if len(tails) == 1 else "_".join(sorted(tails))


def rich(par, text: str, size: float | None = None) -> None:
    """Add `text` to a paragraph, rendering **bold** and *italic* spans."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            r = par.add_run(tok)
        if size:
            r.font.size = Pt(size)


def heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), NAVY


def figure(doc, e: dict, first: bool) -> None:
    png = FIGDIR / f"{e['key']}.png"
    with Image.open(png) as im:
        w_px, h_px = im.size
        dpi = im.info.get("dpi", (600, 600))[0]
    width = min(TEXT_W, Mm(w_px / dpi * 25.4))
    # the figure, its caption and its legend share one page: the figure gives up only the height the
    # legend needs beyond the default allowance
    max_h = min(MAX_FIG_H, PAGE_H - 2 * MARGIN - Mm(CAPTION_MM) - legend_height(e["legend"]))
    if width * h_px / w_px > max_h:
        width = int(max_h * w_px / h_px)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # break on the figure itself: a separate page-break paragraph lands on a page of its own (a blank
    # page) whenever the previous legend fills its page exactly
    pic.paragraph_format.page_break_before = not first
    pic.add_run().add_picture(str(png), width=width)
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(4)
    r = cap.add_run(f"{e['label']}  |  {e['title']}")
    r.bold = True
    leg = doc.add_paragraph()
    leg.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    leg.paragraph_format.space_after = Pt(10)
    rich(leg, e["legend"], size=LEGEND_PT)


def generated_summary(N: dict) -> str:
    p = N['provenance']
    runs = p['runs']
    total = sum(r['n_candidates'] for r in runs.values())
    return (f"This report records {total} candidate transcription factors across the canonical "
            f"benchmark and survey runs, produced with predictor {p['predictor_commit']}. "
            "Benchmark results compare predictions with curated reference evidence; survey results "
            "describe hypotheses without an operator truth set. Operator sites, inducer calls and "
            "regulons remain computational predictions for experimental prioritisation. "
            "All counts, figures and legends are generated from the recorded run statistics.")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pdf', action='store_true', help='require DOCX to PDF conversion')
    parser.add_argument('--out', type=Path, help='write a fresh report artifact at this path')
    args = parser.parse_args(argv)
    N = json.loads((HERE / "numbers.json").read_text(encoding="utf-8"))
    doc = Document()
    # python-docx's template omits this required OOXML zoom attribute.
    for zoom in doc.settings.element.xpath("./w:zoom"):
        zoom.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent", "100")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = MARGIN
    base = doc.styles["Normal"]
    base.font.name = "Arial"
    base.font.size = Pt(10.5)
    base.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Arial")

    p = doc.add_paragraph()
    r = p.add_run("Metalloregulator operator-prediction pipeline")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(17), NAVY
    p = doc.add_paragraph()
    r = p.add_run("Validation and four-genome survey: figures and legends")
    r.font.size, r.font.color.rgb = Pt(13), BLUE
    p = doc.add_paragraph()
    r = p.add_run(provenance_line(N) + f" Run dates: {N['provenance'].get('run_dates', 'n/a')}.")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), GREY

    heading(doc, "Validation summary")
    p = doc.add_paragraph()
    rich(p, generated_summary(N))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    rich(p, "**Definitions used throughout**")
    for key, term in DEFINITIONS:
        d = N["definitions"].get(key)
        if not d:
            continue
        item = doc.add_paragraph(style="List Bullet")
        rich(item, f"**{term}.** {d}", size=9.5)

    entries = figure_entries(N)
    for group, is_main in (("Main figures", True), ("Supplementary figures", False)):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        heading(doc, group)
        for i, e in enumerate(x for x in entries if x["main"] == is_main):
            figure(doc, e, first=i == 0)

    out = args.out or HERE / f"run_report_{run_suffix(N)}_fig4_structures.docx"
    out = out.expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    if args.pdf:
        from pdf_export import export_pdf
        export_pdf(out, required=True)
    print(f"  wrote {out.name} ({len(entries)} figures)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
