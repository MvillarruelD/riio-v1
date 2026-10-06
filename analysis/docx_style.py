#!/usr/bin/env python
"""docx_style.py -- the house style for the report documents: fonts, figures, tables, captions.

Extracted from the original report builder so the style is shared without carrying that report's
body with it. Nothing here reads a run directory or decides a number; callers pass in the figure
path and the rows. That separation is the point -- a style helper that also knows where the data
lives is a style helper that can silently render the wrong run.
"""
from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x60, 0x60, 0x60)
MAXW, MAXH = 6.4, 8.0          # inches of usable space for a figure


def fig_size(path):
    from PIL import Image
    w, h = Image.open(path).size
    ar = h / w
    width = MAXW
    if width * ar > MAXH:
        width = MAXH / ar
    return Inches(width), Inches(width * ar)


# --------------------------------------------------------------------------- doc helpers
def style_doc(doc):
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10.5)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(5)
    for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
        s = doc.styles[f"Heading {lvl}"]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = INK
        s.font.bold = True
        s.paragraph_format.space_before = Pt(14 if lvl < 3 else 10)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True


def para(doc, text="", *, size=10.5, bold=False, italic=False, align=None, color=None,
         space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
    return p


def rich(doc, parts, *, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6):
    """parts = [(text, {'b':..,'i':..}), ...]"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, fmt in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = fmt.get("b", False)
        r.italic = fmt.get("i", False)
        if fmt.get("grey"):
            r.font.color.rgb = GREY
    return p


def body(doc, text, size=10.5):
    return rich(doc, [(text, {})], size=size)


def figure(doc, png, number, title, description, methods):
    """Embed a figure with a structured legend: title, description, methods.

    A MISSING image leaves a visible placeholder and keeps the legend, rather than raising. Two
    reasons, and the second is the important one: an unbuildable figure should not abort a report
    that is otherwise complete (F24 splices in the manuscript's own artwork, so it is absent whenever
    that is), and a figure that simply VANISHED would leave the reader with a numbering gap and no
    idea whether it was dropped on purpose. The placeholder says which file is missing.
    """
    png = Path(png)
    if not png.is_file():
        warn = doc.add_paragraph()
        warn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        warn.paragraph_format.space_before = Pt(10)
        warn.paragraph_format.space_after = Pt(4)
        r = warn.add_run(f"[Figure {number} not generated in this run — {png.name} is absent]")
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        w, h = fig_size(png)
        p.add_run().add_picture(str(png), width=w, height=h)

    leg = doc.add_paragraph()
    leg.alignment = WD_ALIGN_PARAGRAPH.LEFT
    leg.paragraph_format.space_after = Pt(3)
    r = leg.add_run(f"Figure {number}. {title}")
    r.bold = True
    r.font.size = Pt(9)
    r2 = leg.add_run(" " + description)
    r2.font.size = Pt(9)

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.LEFT
    m.paragraph_format.space_after = Pt(12)
    rm = m.add_run("Methods. ")
    rm.bold = True
    rm.italic = True
    rm.font.size = Pt(9)
    rm2 = m.add_run(methods)
    rm2.font.size = Pt(9)
    rm2.italic = False


def table(doc, headers, rows, widths=None, mono_cols=(), size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    hdr = t.rows[0].cells
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(htxt)
        r.bold = True
        r.font.size = Pt(size)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "DCE5E1")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
            if i in mono_cols:
                r.font.name = "Consolas"
            if row_index % 2:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "F5F7F6")
                cells[i]._tc.get_or_add_tcPr().append(shd)
    for row in t.rows:
        for cell in row.cells:
            tc_mar = OxmlElement("w:tcMar")
            for edge, value in (("top", "70"), ("left", "90"),
                                ("bottom", "70"), ("right", "90")):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), value)
                margin.set(qn("w:type"), "dxa")
                tc_mar.append(margin)
            cell._tc.get_or_add_tcPr().append(tc_mar)
    if widths:
        for i, wd in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(wd)
    return t


def tbl_caption(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"Table {n}. ")
    r.bold = True
    r.font.size = Pt(9)
    r2 = p.add_run(text)
    r2.font.size = Pt(9)
    return p


def page_break(doc):
    doc.add_page_break()


# --------------------------------------------------------------------------- content
