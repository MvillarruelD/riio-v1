"""Rebuild portable run reports from existing inputs and bundles, without prediction.

All output is beside the bundles. Missing/failed candidates remain table rows;
only existing REPORT.html files receive hyperlinks.
"""
from __future__ import annotations

import csv
import html
import json
from pathlib import Path
from urllib.parse import quote

from predictor.report import discovery_page, run_page


def _json(path):
    if not path.is_file():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _object(path):
    value = _json(path)
    return value if isinstance(value, dict) else {}


def _table(path):
    if not path.is_file():
        return []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t" if path.suffix == ".tsv" else ","))


def _manifest(run, manifest_path=None):
    if manifest_path:
        return _table(Path(manifest_path))
    for name in ("manifest.csv", "manifest.tsv", "discovery/source_manifest.csv"):
        if (run / name).is_file():
            return _table(run / name)
    return []


def _stored_bundles(run):
    """Return completed visible children, excluding interrupted publication drafts and escapes."""
    for parent in (run / "jobs", run):
        bundles = {}
        for path in parent.glob("*/dossier.json"):
            if path.parent.name.startswith(".") or not path.resolve().is_relative_to(run.resolve()):
                continue
            bundles[path.parent.name] = path.parent
        if bundles:
            return bundles
    return {}


def _validate_run(run, manifest_path=None):
    if not run.is_dir():
        raise ValueError(f"run directory does not exist: {run}")
    if (run / "dossier.json").is_file():
        raise ValueError("report destination must be outside an existing TF bundle")
    if manifest_path and not Path(manifest_path).is_file():
        raise ValueError(f"source manifest does not exist: {manifest_path}")
    if not manifest_path and not any((run / name).is_file() for name in ("manifest.csv", "manifest.tsv", "summary.tsv")) and not _stored_bundles(run) and not any((run / "discovery").glob("candidates.*")):
        raise ValueError("directory contains no stored run inputs, candidates, or TF bundles")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _escape(value):
    return html.escape(_text(value), quote=True)


def _href(relative):
    return quote(str(relative).replace("\\", "/"), safe="/")


def _safe_file(run, relative):
    """Resolve untrusted manifest paths only inside the run."""
    target = (run / str(relative)).resolve()
    return target if target.is_relative_to(run) and target.is_file() else None


def _write_tsv(path, rows, columns):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, columns, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        writer.writerows({k: _text(v) for k, v in row.items()} for row in rows)


CSS = """body{font:15px Arial,'Liberation Sans',sans-serif;color:#20313b;background:#f7f9f9;margin:0}
main{max-width:1400px;margin:auto;padding:36px}h1{font-size:32px}h2{margin-top:32px}
a{color:#176b79}p{line-height:1.6}table{border-collapse:collapse;background:white;width:100%;font-size:13px}
td,th{padding:10px;text-align:left;border-bottom:1px solid #dce5e7;vertical-align:top;overflow-wrap:anywhere}
th{background:#e9f0f2;cursor:pointer}button{border:0;background:transparent;font:inherit;font-weight:bold;cursor:pointer}
.scroll{overflow:auto}.counts{display:flex;gap:12px;flex-wrap:wrap}.count{background:white;padding:18px;min-width:130px}
.bar{height:8px;background:#4c8790;margin:5px 0}.muted{color:#637780}pre{white-space:pre-wrap;overflow-wrap:anywhere}
"""
SORT = """<script>
document.querySelectorAll('th button').forEach(b=>b.addEventListener('click',()=>{
const h=b.parentElement,t=h.closest('table'),i=Array.from(h.parentElement.children).indexOf(h);
const asc=h.getAttribute('aria-sort')!=='ascending';
t.querySelectorAll('thead th').forEach(x=>x.setAttribute('aria-sort','none'));
h.setAttribute('aria-sort',asc?'ascending':'descending');
t.querySelectorAll('tbody').forEach(body=>{
const rows=Array.from(body.rows).filter(r=>!r.classList.contains('group'));
rows.sort((a,b)=>{const value=cell=>cell.getAttribute('data-sort-value')??cell.textContent.trim();
const x=value(a.cells[i]),y=value(b.cells[i]);
const n=(x!==''&&y!==''&&Number.isFinite(Number(x))&&Number.isFinite(Number(y)))?Number(x)-Number(y):x.localeCompare(y);
return asc?n:-n;});rows.forEach(r=>body.appendChild(r));});}));
</script>"""


def _page(title, body, *, css=""):
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{_escape(title)}</title><style>{CSS}{css}</style></head><body><main>'
            f'<h1>{_escape(title)}</h1>{body}</main>{SORT}</body></html>')


def write_discovery_report(run_dir, manifest_path=None):
    """Write discovery HTML/TSV/FASTA, preserving available raw engine evidence."""
    run = Path(run_dir).resolve()
    _validate_run(run, manifest_path)
    discovery = run / "discovery"
    discovery.mkdir(parents=True, exist_ok=True)
    if manifest_path:
        import shutil
        source = Path(manifest_path).resolve()
        target = discovery / "source_manifest.csv"
        if source != target.resolve():
            shutil.copyfile(source, target)
    manifest = _manifest(run, manifest_path)
    raw_path = discovery / "candidates.tsv"
    raw = _table(raw_path)
    retained = _json(discovery / "candidates.json")
    if isinstance(retained, list):
        raw = [{**(r.get("evidence") or {}), **r} for r in retained if isinstance(r, dict)] + raw
    # BITACORA's own candidate export is raw evidence; never overwrite it.
    raw_index = {}
    for record in raw:
        raw_index.setdefault(record.get("protein_id", record.get("locus_tag", "")), record)
    columns = ["run_name", "family", "clade", "contig", "start", "end", "strand", "evalue", "bits",
               "discovery_engine", "discovery_profile", "trim_start", "trim_end", "evidence"]
    rows = []
    for item in manifest:
        name = item.get("run_name", item.get("name", ""))
        pid = item.get("locus_tag") or item.get("protein_id") or name.split("__")[-1]
        fasta = _safe_file(run, item.get("fasta", ""))
        if fasta:
            first = next(iter(fasta.read_text(encoding="utf-8").splitlines()), "")
            if first.startswith(">"):
                pid = first[1:].split()[0]
        evidence = raw_index.get(pid, {})
        row = {**evidence, **item, "run_name": name,
               "clade": item.get("clade") or evidence.get("clade") or evidence.get("profile") or item.get("discovery_profile", ""),
               "evalue": item.get("evalue") or evidence.get("evalue", ""),
               "bits": item.get("bits") or evidence.get("bits", ""),
               "evidence": evidence or item.get("evidence") or "not retained in legacy manifest"}
        rows.append(row)
    if not rows:
        rows = raw
    if not rows:
        # Older canonical runs have bundles but no run manifest. Expose their
        # recovered candidate identities without claiming retained discovery evidence.
        for bundle in sorted(_stored_bundles(run).values()):
            path = bundle / "dossier.json"
            dossier = _object(path)
            locus = dossier.get("tf_locus") or ["", "", ""]
            rows.append({"run_name": path.parent.name, "family": dossier.get("family", "unknown"),
                         "clade": dossier.get("ssn_cluster", ""), "contig": dossier.get("genome_accession", ""),
                         "start": locus[0] if len(locus) > 0 else "", "end": locus[1] if len(locus) > 1 else "",
                         "strand": locus[2] if len(locus) > 2 else "", "discovery_engine": "not recorded",
                         "evidence": "identity recovered from existing prediction bundle; discovery evidence unavailable"})
    _write_tsv(discovery / "discovery_candidates.tsv", rows, columns)
    # If no raw candidate table exists, provide the requested canonical name too.
    if not raw_path.exists():
        _write_tsv(raw_path, rows, columns)
    sequences = []
    for item in manifest:
        fasta = _safe_file(run, item.get("fasta", ""))
        if fasta:
            sequences.append(fasta.read_text(encoding="utf-8").rstrip() + "\n")
    if sequences:
        (discovery / "candidates.faa").write_text("".join(sequences), encoding="utf-8")
    # Per-TF report links, relative to discovery/, for candidates that were predicted.
    reports = {name: _href(Path("..", bundle.relative_to(run), "REPORT.html").as_posix())
               for name, bundle in _stored_bundles(run).items()
               if _safe_file(run, (bundle / "REPORT.html").relative_to(run))}
    body = discovery_page.render(rows, discovery=discovery, reports=reports)
    target = discovery / "DISCOVERY_REPORT.html"
    target.write_text(_page("Discovery report", body, css=run_page.CSS + discovery_page.CSS), encoding="utf-8")
    return target


def _run_rows(run, manifest_path=None):
    from predictor.effector.inducer_vocab import inducer_class_of, is_abstention
    manifests = _manifest(run, manifest_path)
    summary = {r.get("run_name", r.get("name", "")): r for r in _table(run / "summary.tsv")}
    bundles = _stored_bundles(run)
    names = list(dict.fromkeys([r.get("run_name", r.get("name", "")) for r in manifests] + list(summary) + list(bundles)))
    manifest_index = {r.get("run_name", r.get("name", "")): r for r in manifests}
    result = []
    for name in names:
        if not name:
            continue
        item = manifest_index.get(name, {})
        row = {**item, **summary.get(name, {}), "run_name": name}
        bundle = bundles.get(name)
        dossier = _object(bundle / "dossier.json") if bundle else {}
        bundle_manifest = _object(bundle / "bundle_manifest.json") if bundle else {}
        inducer = dossier.get("inducers") or {}
        top = inducer.get("top") or dossier.get("effector") or row.get("inducer_top", "")
        operators = dossier.get("operators") or {}
        primary = operators.get("primary") or {} if isinstance(operators, dict) else {}
        provenance = primary.get("provenance") or {}
        run_status = row.get("status", "")
        if dossier and not run_status.startswith(("failed", "timed out")):
            run_status = bundle_manifest.get("status") or "bundle status not recorded"
        # A candidate without a bundle was never predicted: it is not an abstention, and its only
        # family is the discovery routing hint (summary.tsv writes an empty family for such rows).
        predicted = bool(dossier)
        row.update(family=(dossier.get("family") or "unknown") if predicted else "",
                   discovery_family=item.get("family") or row.get("discovery_family", ""),
                   ssn_cluster=dossier.get("ssn_cluster", row.get("ssn_cluster", "")),
                   inducer_top=top,
                   inducer_class=(inducer_class_of(top) or "unresolved") if predicted else "not predicted",
                   agreement=inducer.get("agreement", row.get("agreement", "")),
                   primary_operator=primary.get("sequence", row.get("primary_operator", "")),
                   primary_start=provenance.get("start", row.get("primary_start", "")),
                   primary_end=provenance.get("end", row.get("primary_end", "")),
                   # Bundles record where the rank-1 site sits, not a confidence score.
                   primary_locality=provenance.get("locality") or row.get("primary_locality", ""),
                   operator_confidence=primary.get("confidence"),
                   abstention=is_abstention(top) if predicted else "",
                   n_regulon_operons=len(dossier.get("regulon") or []) if dossier else row.get("n_regulon_operons", ""),
                   status=run_status or "not run / bundle missing")
        report = _safe_file(run, (bundle / "REPORT.html").relative_to(run)) if bundle else None
        row["report_path"] = report.relative_to(run).as_posix() if report and report.is_file() else ""
        if dossier and not row["report_path"] and not row["status"].startswith(("failed", "timed out")):
            row["status"] = "bundle present / report missing"
        if bundle and not dossier:
            row["status"] = "bundle present / dossier missing or invalid"
        # Underscore keys feed the HTML page only; the S1 tables never include them.
        record = _safe_file(run, (bundle / "run" / "run_record.json").relative_to(run)) if bundle else None
        row.update(_predicted=predicted, _dossier=dossier, _bundle_manifest=bundle_manifest,
                   _run_record=_object(record) if record else {},
                   _href=_href(row["report_path"]) if row["report_path"] else "",
                   _record_href=_href(record.relative_to(run).as_posix()) if record else "")
        result.append(row)
    return result, bundles


def _write_xlsx(path, rows, columns):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TF summary"
    sheet.append(columns)
    for row in rows:
        # Force source text to literal cells; labels beginning '=' are not formulas.
        sheet.append([_text(row.get(column, "")) for column in columns])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
        link = row.get("report_path")
        if link:
            cell = sheet.cell(sheet.max_row, columns.index("report_path") + 1)
            cell.hyperlink = _href(link)
            cell.style = "Hyperlink"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="264653")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 42 if column in ("product", "report_path", "primary_operator") else 22
    workbook.save(path)


def write_run_report(run_dir, manifest_path=None):
    """Write global HTML and linked XLSX/TSV from stored run data only."""
    run = Path(run_dir).resolve()
    _validate_run(run, manifest_path)
    rows, _ = _run_rows(run, manifest_path)
    columns = ["run_name", "locus_tag", "gene", "product", "discovery_family", "family", "ssn_cluster",
               "inducer_top", "inducer_class", "agreement", "primary_operator", "primary_start", "primary_end",
               "primary_locality", "operator_confidence", "n_regulon_operons", "abstention", "status", "report_path"]
    columns += sorted({key for row in rows for key in row if not key.startswith("_")} - set(columns))
    _write_tsv(run / "Supplementary_Table_S1.tsv", rows, columns)
    _write_xlsx(run / "Supplementary_Table_S1.xlsx", rows, columns)
    links = ['<a href="Supplementary_Table_S1.xlsx">Supplementary Table S1 (XLSX)</a>',
             '<a href="Supplementary_Table_S1.tsv">TSV</a>']
    for relative, label in (("discovery/DISCOVERY_REPORT.html", "Discovery report"), ("af3/af3_manifest.tsv", "AF3 manifest")):
        if (run / relative).is_file():
            links.append(f'<a href="{relative}">{label}</a>')
    export_warnings = _json(run / "af3" / "export_warnings.json")
    if export_warnings:
        links.append('<a href="af3/export_warnings.json">AF3 export warnings</a>')
    body = run_page.render(rows, run_input=_object(run / "run_input.json"), links=links,
                           af3_warnings=export_warnings)
    target = run / "RUN_REPORT.html"
    target.write_text(_page("RIIO run report", body, css=run_page.CSS), encoding="utf-8")
    return target
